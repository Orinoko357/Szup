from __future__ import annotations
import csv
import io
from typing import Any

from fastapi.responses import StreamingResponse


def send_xlsx(data: list[dict], cols: list[dict], filename: str, sheet_name: str = "Sheet1") -> StreamingResponse:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise ImportError("openpyxl is not installed")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for i, col in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i, value=col["header"])
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row in enumerate(data, 2):
        for col_idx, col in enumerate(cols, 1):
            value = row.get(col["key"])
            if value is not None and hasattr(value, "isoformat"):
                value = str(value)
            ws.cell(row=row_idx, column=col_idx, value=value)

    for col in ws.column_dimensions:
        ws.column_dimensions[col].width = 18

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}.xlsx"},
    )


def send_csv(data: list[dict], cols: list[dict], filename: str) -> StreamingResponse:
    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")
    writer.writerow([col["header"] for col in cols])
    for row in data:
        writer.writerow([row.get(col["key"], "") for col in cols])
    output.seek(0)

    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={filename}.csv"},
    )
